import hashlib

import pytest
from app.models import Documento, DocumentoChunk, DocumentoPagina
from app.services import document_processor, ingestion_service
from app.services.document_eligibility import (
    condicoes_consulta,
    condicoes_modelo_minuta,
    documento_elegivel_consulta,
    modelo_minuta_elegivel,
)
from app.services.document_processor import (
    extrair_documento,
)
from docx import Document
from docx.oxml import OxmlElement
from pypdf import PdfWriter
from pypdf.generic import (
    DecodedStreamObject,
    DictionaryObject,
    NameObject,
    NumberObject,
)


def criar_pdf(caminho, tipos):
    writer = PdfWriter()
    for tipo in tipos:
        pagina = writer.add_blank_page(width=300, height=400)
        if tipo == "vazia":
            continue
        recursos = DictionaryObject()
        stream = DecodedStreamObject()
        if tipo == "imagem":
            imagem = DecodedStreamObject()
            imagem.set_data(b"\x00")
            imagem.update(
                {
                    NameObject("/Type"): NameObject("/XObject"),
                    NameObject("/Subtype"): NameObject("/Image"),
                    NameObject("/Width"): NumberObject(1),
                    NameObject("/Height"): NumberObject(1),
                    NameObject("/ColorSpace"): NameObject("/DeviceGray"),
                    NameObject("/BitsPerComponent"): NumberObject(8),
                }
            )
            recursos[NameObject("/XObject")] = DictionaryObject(
                {NameObject("/Img"): writer._add_object(imagem)}
            )
            stream.set_data(b"q 200 0 0 200 0 0 cm /Img Do Q")
        else:
            fonte = DictionaryObject(
                {
                    NameObject("/Type"): NameObject("/Font"),
                    NameObject("/Subtype"): NameObject("/Type1"),
                    NameObject("/BaseFont"): NameObject("/Helvetica"),
                }
            )
            recursos[NameObject("/Font")] = DictionaryObject(
                {NameObject("/F1"): writer._add_object(fonte)}
            )
            stream.set_data(
                b"BT /F1 12 Tf 10 200 Td (Regra sintetica: conferir os documentos antes do ato.) Tj ET"
            )
        pagina[NameObject("/Resources")] = recursos
        pagina[NameObject("/Contents")] = writer._add_object(stream)
    writer.write(caminho)


def test_docx_preserva_tabelas_aninhadas_mesclas_e_ordem(tmp_path):
    caminho = tmp_path / "tabelas.docx"
    doc = Document()
    doc.add_paragraph("INICIO")
    tabela = doc.add_table(rows=1, cols=2)
    tabela.cell(0, 0).merge(tabela.cell(0, 1)).text = "CELULA_MESCLADA"
    interna = tabela.cell(0, 0).add_table(rows=1, cols=1)
    interna.cell(0, 0).text = "CONTEUDO_ANINHADO"
    doc.add_paragraph("FINAL")
    doc.save(caminho)
    original = hashlib.sha256(caminho.read_bytes()).digest()
    resultado = extrair_documento(str(caminho))
    texto = "\n\n".join(parte.conteudo for parte in resultado.partes)
    assert texto.count("CELULA_MESCLADA") == 1
    assert (
        texto.index("INICIO")
        < texto.index("CELULA_MESCLADA")
        < texto.index("CONTEUDO_ANINHADO")
        < texto.index("FINAL")
    )
    assert resultado.situacao == "PROCESSADO_COMPLETO"
    assert hashlib.sha256(caminho.read_bytes()).digest() == original


def test_docx_grande_em_tabela_nao_perde_corpo(tmp_path):
    caminho = tmp_path / "grande.docx"
    doc = Document()
    doc.add_paragraph("Titulo")
    tabela = doc.add_table(rows=100, cols=1)
    for i, linha in enumerate(tabela.rows):
        linha.cells[0].text = f"REGRA_{i:03d} " + "Conteudo de avaliacao. " * 40
    doc.save(caminho)
    resultado = extrair_documento(str(caminho))
    texto = "\n\n".join(parte.conteudo for parte in resultado.partes)
    chunks = document_processor.limpar_texto_extraido(texto)
    from app.services.chunking import dividir_texto

    assert len(dividir_texto(chunks)) > 1
    assert all(f"REGRA_{i:03d}" in texto for i in range(100))
    assert len(texto) > 80000
    assert len(resultado.partes) > 1


def test_docx_imagem_sinaliza_parcial(tmp_path):
    caminho = tmp_path / "imagem.docx"
    doc = Document()
    doc.add_paragraph("Texto legivel")
    doc.add_paragraph().add_run()._r.append(OxmlElement("w:drawing"))
    doc.save(caminho)
    resultado = extrair_documento(str(caminho))
    assert resultado.situacao == "EXTRACAO_PARCIAL"
    assert resultado.avisos


@pytest.mark.parametrize(
    "tipos,esperado",
    [
        (["texto", "texto"], "PROCESSADO_COMPLETO"),
        (["texto", "imagem", "texto"], "EXTRACAO_PARCIAL"),
        (["imagem", "imagem"], "NECESSITA_OCR"),
        (["vazia"], "SEM_TEXTO"),
        (["texto", "vazia"], "PROCESSADO_COMPLETO"),
    ],
)
def test_pdf_preserva_paginas_e_diagnostica_sem_ocr(tmp_path, tipos, esperado):
    caminho = tmp_path / "amostra.pdf"
    criar_pdf(caminho, tipos)
    resultado = extrair_documento(str(caminho))
    assert len(resultado.partes) == len(tipos)
    assert [p.numero for p in resultado.partes] == list(range(1, len(tipos) + 1))
    assert resultado.situacao == esperado
    assert resultado.resumo()["partes_necessitam_ocr"] == tipos.count("imagem")
    assert resultado.resumo()["partes_vazias"] == tipos.count("vazia")


def test_pdf_permite_forcar_ocr_integral_sem_alterar_original(tmp_path, monkeypatch):
    caminho = tmp_path / "texto.pdf"
    criar_pdf(caminho, ["texto", "texto"])
    original = caminho.read_bytes()
    situacoes = []

    def registrar(_caminho, extracao):
        situacoes.extend(parte.situacao for parte in extracao.partes)
        return extracao

    monkeypatch.setattr(document_processor, "aplicar_ocr", registrar)
    extrair_documento(str(caminho), forcar_ocr=True)
    assert situacoes == ["NECESSITA_OCR", "NECESSITA_OCR"]
    assert caminho.read_bytes() == original


@pytest.mark.parametrize(
    "estado",
    [
        "PENDENTE_VERIFICACAO",
        "EXTRACAO_PARCIAL",
        "NECESSITA_OCR",
        "SEM_TEXTO",
        "ERRO_PROCESSAMENTO",
    ],
)
def test_extracao_incompleta_fora_dos_dois_acervos(db, usuario_factory, estado):
    usuario = usuario_factory("integridade")
    norma = Documento(
        titulo="Sintetico",
        tipo="NORMA",
        situacao="APROVADO",
        status="PRONTO",
        status_seguranca="LIBERADO",
        situacao_extracao=estado,
        nome_arquivo="a.txt",
        caminho_arquivo="a.txt",
        criado_por=usuario.id,
    )
    db.add(norma)
    db.commit()
    assert not documento_elegivel_consulta(norma)
    assert db.query(Documento).filter(*condicoes_consulta()).count() == 0
    norma.tipo = "MODELO_MINUTA"
    norma.tipo_ato = "COMPRA_VENDA"
    db.commit()
    assert not modelo_minuta_elegivel(norma, "COMPRA_VENDA")
    assert db.query(Documento).filter(*condicoes_modelo_minuta()).count() == 0


def test_ingestao_parcial_preserva_pagina_sem_texto_e_governanca(
    db, testing_session_factory, usuario_factory, tmp_path, monkeypatch
):
    usuario = usuario_factory("pagina-parcial")
    arquivo = tmp_path / "misto.pdf"
    criar_pdf(arquivo, ["texto", "imagem", "texto"])
    original = arquivo.read_bytes()
    documento = Documento(
        titulo="Sintetico",
        tipo="NORMA",
        situacao="APROVADO",
        nome_arquivo=arquivo.name,
        caminho_arquivo=str(arquivo),
        criado_por=usuario.id,
    )
    db.add(documento)
    db.commit()
    monkeypatch.setattr(ingestion_service, "SessionLocal", testing_session_factory)
    monkeypatch.setattr(
        ingestion_service, "gerar_embeddings_documento", lambda *args, **kwargs: 2
    )
    ingestion_service.processar_documento(documento.id)
    db.expire_all()
    assert documento.situacao == "APROVADO"
    assert documento.situacao_extracao == "EXTRACAO_PARCIAL"
    assert documento.total_paginas == 3
    assert documento.diagnostico_extracao["partes_originais"] == 3
    assert db.query(DocumentoPagina).filter_by(documento_id=documento.id).count() == 3
    assert (
        db.query(DocumentoPagina)
        .filter_by(documento_id=documento.id, pagina=2)
        .one()
        .situacao_extracao
        == "NECESSITA_OCR"
    )
    assert not documento_elegivel_consulta(documento)
    ingestion_service.processar_documento(documento.id)
    db.expire_all()
    assert db.query(DocumentoPagina).filter_by(documento_id=documento.id).count() == 3
    assert (
        db.query(DocumentoChunk).filter_by(documento_id=documento.id).count()
        == documento.total_chunks
    )
    assert arquivo.read_bytes() == original


def test_filtro_extracao_e_diagnostico_na_api(
    client, db, usuario_factory, auth_headers
):
    usuario = usuario_factory("filtro-extracao")
    for i in range(21):
        db.add(
            Documento(
                titulo=f"Sintetico {i}",
                tipo="NORMA",
                nome_arquivo="a.txt",
                caminho_arquivo="caminho-privado",
                criado_por=usuario.id,
                situacao_extracao="EXTRACAO_PARCIAL",
            )
        )
    db.commit()
    resposta = client.get(
        "/api/documentos?situacao_extracao=EXTRACAO_PARCIAL&pagina=2&por_pagina=20",
        headers=auth_headers(usuario),
    )
    assert resposta.status_code == 200
    dados = resposta.json()
    assert dados["total"] == 21 and len(dados["items"]) == 1
    assert dados["items"][0]["situacao_extracao"] == "EXTRACAO_PARCIAL"
    assert "caminho-privado" not in resposta.text


def test_falha_de_embedding_preserva_derivados_anteriores(
    db, testing_session_factory, usuario_factory, tmp_path, monkeypatch
):
    usuario = usuario_factory("falha-derivados")
    arquivo = tmp_path / "regra.txt"
    arquivo.write_text("Regra sintetica atualizada.", encoding="utf-8")
    documento = Documento(
        titulo="Sintetico",
        tipo="NORMA",
        nome_arquivo=arquivo.name,
        caminho_arquivo=str(arquivo),
        criado_por=usuario.id,
        status="PRONTO",
        situacao="APROVADO",
        situacao_extracao="PROCESSADO_COMPLETO",
        total_paginas=1,
        total_chunks=1,
    )
    db.add(documento)
    db.flush()
    pagina = DocumentoPagina(
        documento_id=documento.id, pagina=1, conteudo="Versao anterior."
    )
    chunk = DocumentoChunk(
        documento_id=documento.id, pagina=1, posicao=1, conteudo="Versao anterior."
    )
    db.add_all([pagina, chunk])
    db.commit()
    ids_anteriores = (pagina.id, chunk.id)
    monkeypatch.setattr(ingestion_service, "SessionLocal", testing_session_factory)

    def falhar(*args, **kwargs):
        raise RuntimeError("detalhe privado que nao deve sair na API")

    monkeypatch.setattr(ingestion_service, "gerar_embeddings_documento", falhar)
    ingestion_service.processar_documento(documento.id)
    db.expire_all()
    assert documento.situacao_extracao == "ERRO_PROCESSAMENTO"
    assert documento.situacao == "APROVADO"
    assert db.get(DocumentoPagina, ids_anteriores[0]).conteudo == "Versao anterior."
    assert db.get(DocumentoChunk, ids_anteriores[1]).conteudo == "Versao anterior."
    assert "privado" not in documento.erro_processamento
    assert arquivo.read_text(encoding="utf-8") == "Regra sintetica atualizada."


def test_diagnostico_nao_grava_nem_expoe_texto(db, usuario_factory, tmp_path):
    from scripts.diagnose_extraction import diagnosticar

    usuario = usuario_factory("diagnostico-leitura")
    arquivo = tmp_path / "teste.txt"
    arquivo.write_text("CONTEUDO_PRIVADO_TESTE", encoding="utf-8")
    documento = Documento(
        titulo="TITULO_PRIVADO",
        tipo="NORMA",
        nome_arquivo=arquivo.name,
        caminho_arquivo=str(arquivo),
        criado_por=usuario.id,
    )
    db.add(documento)
    db.commit()
    resultado = diagnosticar(db, documento.id)
    assert resultado["chunks_recalculados_em_memoria"] == 1
    assert "PRIVADO" not in str(resultado)
    assert documento.situacao_extracao == "PENDENTE_VERIFICACAO"
    assert db.query(DocumentoPagina).count() == 0
    assert db.query(DocumentoChunk).count() == 0

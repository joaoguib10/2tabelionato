import json
import logging
import re
from io import BytesIO
from pathlib import Path
from uuid import UUID

from app.auth import get_current_user
from app.config import MINUTA_MODEL_CONTEXT_LIMIT, MINUTA_TOTAL_CONTEXT_LIMIT
from app.dependencies import get_db
from app.models import Documento, DocumentoPagina, Usuario
from app.permissions import require_roles
from app.schemas import (
    MinutaAnaliseResponse,
    MinutaExportRequest,
    MinutaResponse,
    ModeloMinutaResponse,
    ModeloUtilizadoResponse,
)
from app.services.document_eligibility import (
    TIPOS_ATO,
    condicoes_modelo_minuta,
    modelo_minuta_elegivel,
)
from app.services.minute_service import (
    analisar_arquivo,
    estruturar_analise,
    gerar_minuta,
)
from docx import Document as DocxDocument
from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile, status
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session

logger = logging.getLogger(__name__)
router = APIRouter(
    prefix="/api/minutas",
    tags=["Minutas"],
    dependencies=[Depends(require_roles("ADMIN"))],
)
EXTENSOES = (".pdf", ".jpg", ".jpeg", ".png")
LIMITE_ARQUIVO = 15 * 1024 * 1024
LIMITE_TOTAL = 50 * 1024 * 1024
LIMITE_TEXTO_ARQUIVO = 25_000


def _carregar_json(valor: str, campo: str):
    try:
        return json.loads(valor)
    except json.JSONDecodeError as erro:
        raise HTTPException(
            status_code=422, detail=f"O campo {campo} não contém JSON válido."
        ) from erro


async def _analisar_uploads(
    grupos: list[tuple[str, list[UploadFile] | None]],
) -> tuple[list[dict], list[str]]:
    resultados: list[dict] = []
    avisos: list[str] = []
    total = 0
    for finalidade, arquivos in grupos:
        for arquivo in arquivos or []:
            nome = Path(arquivo.filename or "arquivo").name or "arquivo"
            if not nome.lower().endswith(EXTENSOES):
                raise HTTPException(
                    status_code=400, detail=f"Formato não permitido em {nome}."
                )
            conteudo = await arquivo.read()
            total += len(conteudo)
            if not conteudo or len(conteudo) > LIMITE_ARQUIVO or total > LIMITE_TOTAL:
                raise HTTPException(
                    status_code=413,
                    detail="Use até 15 MB por arquivo e 50 MB no total.",
                )
            assinatura_valida = (
                (nome.lower().endswith(".pdf") and conteudo.startswith(b"%PDF-"))
                or (
                    nome.lower().endswith((".jpg", ".jpeg"))
                    and conteudo.startswith(b"\xff\xd8\xff")
                )
                or (
                    nome.lower().endswith(".png")
                    and conteudo.startswith(b"\x89PNG\r\n\x1a\n")
                )
            )
            if not assinatura_valida:
                raise HTTPException(
                    status_code=400,
                    detail=f"O conteúdo de {nome} não corresponde à extensão.",
                )
            try:
                texto = analisar_arquivo(nome, conteudo, finalidade)
                if texto:
                    texto = texto[:LIMITE_TEXTO_ARQUIVO]
                    resultados.append(
                        {
                            "arquivo": nome,
                            "finalidade": finalidade,
                            "conteudo": texto,
                            "dados_extraidos": estruturar_analise(texto, finalidade),
                        }
                    )
                else:
                    avisos.append(
                        f"{nome}: não há texto extraível; envie as páginas como imagem."
                    )
            except Exception as erro:
                logger.warning(
                    "Falha ao analisar anexo temporário de minuta (tipo=%s).",
                    type(erro).__name__,
                )
                avisos.append(f"{nome}: não foi possível analisar o arquivo.")
    return resultados, avisos


def _obter_modelo(
    db: Session,
    modelo_id: UUID,
    tipo_ato: str,
) -> tuple[Documento, str]:
    modelo = db.get(Documento, modelo_id)
    if modelo is None or not modelo_minuta_elegivel(modelo, tipo_ato):
        raise HTTPException(
            status_code=422,
            detail="O modelo escolhido não está disponível para este tipo de ato.",
        )
    paginas = (
        db.query(DocumentoPagina)
        .filter(DocumentoPagina.documento_id == modelo.id)
        .order_by(DocumentoPagina.pagina.asc())
        .all()
    )
    conteudo = "\n\n".join(pagina.conteudo for pagina in paginas).strip()
    if not conteudo:
        raise HTTPException(
            status_code=409, detail="O modelo aprovado não possui texto processado."
        )
    if len(conteudo) > MINUTA_MODEL_CONTEXT_LIMIT:
        raise HTTPException(
            status_code=413,
            detail=(
                "O modelo aprovado excede o limite de contexto. "
                "Cadastre uma versão mais objetiva antes de gerar a minuta."
            ),
        )
    return modelo, conteudo


def _validar_contexto_total(
    modelo_texto: str,
    dados: dict,
    documentos: list[dict],
) -> None:
    tamanho = sum(
        (
            len(modelo_texto),
            len(json.dumps(dados, ensure_ascii=False)),
            len(json.dumps(documentos, ensure_ascii=False)),
        )
    )
    if tamanho > MINUTA_TOTAL_CONTEXT_LIMIT:
        raise HTTPException(
            status_code=413,
            detail=(
                "O conjunto de modelo, dados e documentos excede o limite de contexto. "
                "Reduza o conteúdo antes de gerar a minuta."
            ),
        )


@router.get("/modelos", response_model=list[ModeloMinutaResponse])
def listar_modelos(
    tipo_ato: str,
    db: Session = Depends(get_db),
    usuario_atual: Usuario = Depends(get_current_user),
):
    ato = tipo_ato.upper()
    if ato not in TIPOS_ATO:
        raise HTTPException(status_code=422, detail="Tipo de ato inválido.")
    modelos = (
        db.query(Documento)
        .filter(*condicoes_modelo_minuta(ato))
        .order_by(Documento.vigencia_inicio.desc(), Documento.updated_at.desc())
        .all()
    )
    return [
        ModeloMinutaResponse(
            id=str(modelo.id),
            titulo=modelo.titulo,
            tipo_ato=modelo.tipo_ato or ato,
            versao=modelo.versao,
            orgao_origem=modelo.orgao_origem,
            vigencia_inicio=modelo.vigencia_inicio,
            updated_at=modelo.updated_at,
        )
        for modelo in modelos
    ]


@router.post("/analisar", response_model=MinutaAnaliseResponse)
async def analisar_documentos_temporarios(
    certidoes_alienantes: list[UploadFile] | None = File(None),
    certidoes_adquirentes: list[UploadFile] | None = File(None),
    matriculas: list[UploadFile] | None = File(None),
    usuario_atual: Usuario = Depends(get_current_user),
):
    analisados, avisos = await _analisar_uploads(
        [
            ("certidão do vendedor/doador", certidoes_alienantes),
            ("certidão do comprador/donatário", certidoes_adquirentes),
            ("matrícula do imóvel", matriculas),
        ]
    )
    return MinutaAnaliseResponse(documentos=analisados, avisos=avisos)


@router.post("/gerar", response_model=MinutaResponse)
async def criar_minuta(
    tipo_ato: str = Form(...),
    modelo_id: UUID = Form(...),
    dados_json: str = Form(...),
    documentos_json: str = Form("[]"),
    confirmado: bool = Form(False),
    certidoes_alienantes: list[UploadFile] | None = File(None),
    certidoes_adquirentes: list[UploadFile] | None = File(None),
    matriculas: list[UploadFile] | None = File(None),
    db: Session = Depends(get_db),
    usuario_atual: Usuario = Depends(get_current_user),
):
    ato = tipo_ato.upper()
    if ato not in TIPOS_ATO:
        raise HTTPException(status_code=422, detail="Tipo de ato inválido.")
    if not confirmado:
        raise HTTPException(
            status_code=422,
            detail="Confirme os dados e documentos antes de gerar a minuta.",
        )
    dados = _carregar_json(dados_json, "dados_json")
    documentos = _carregar_json(documentos_json, "documentos_json")
    if not isinstance(dados, dict) or not isinstance(documentos, list):
        raise HTTPException(status_code=422, detail="Dados de geração inválidos.")
    if len(documentos_json) > 100_000:
        raise HTTPException(
            status_code=413, detail="A análise temporária excede o limite."
        )

    modelo, modelo_texto = _obter_modelo(db, modelo_id, ato)
    _validar_contexto_total(modelo_texto, dados, documentos)
    adicionais, avisos = await _analisar_uploads(
        [
            ("certidão do vendedor/doador", certidoes_alienantes),
            ("certidão do comprador/donatário", certidoes_adquirentes),
            ("matrícula do imóvel", matriculas),
        ]
    )
    documentos.extend(adicionais)
    _validar_contexto_total(modelo_texto, dados, documentos)
    try:
        minuta = gerar_minuta(ato, dados, documentos, modelo_texto)
    except Exception as erro:
        logger.error("Falha local ao gerar minuta (tipo=%s).", type(erro).__name__)
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail="Não foi possível gerar a minuta no serviço local.",
        ) from erro
    avisos.append(
        "Rascunho gerado por IA: confira dados, documentos, tributos e exigências antes do uso."
    )
    modelo_usado = ModeloUtilizadoResponse(
        id=str(modelo.id),
        titulo=modelo.titulo,
        versao=modelo.versao,
        orgao_origem=modelo.orgao_origem,
    )
    return MinutaResponse(
        minuta=minuta,
        avisos=avisos,
        modelo_utilizado=modelo_usado,
    )


@router.post("/exportar")
def exportar_minuta(
    dados: MinutaExportRequest,
    usuario_atual: Usuario = Depends(get_current_user),
):
    documento = DocxDocument()
    documento.add_heading(dados.titulo, level=1)
    for bloco in re.split(r"\n\s*\n", dados.minuta.strip()):
        documento.add_paragraph(bloco.strip())
    documento.add_paragraph("Rascunho sujeito à conferência profissional.")
    arquivo = BytesIO()
    documento.save(arquivo)
    arquivo.seek(0)
    return StreamingResponse(
        arquivo,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": 'attachment; filename="minuta-tabeleao.docx"'},
    )
